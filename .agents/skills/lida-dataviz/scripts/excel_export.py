"""Build an .xlsx with a NATIVE (editable, vector) Excel chart from an aggregated table.

Used by render.py --xlsx. The chart code provides:
    chart_data(data) -> pd.DataFrame   # wide table: 1st column = categories / x, other columns = series
    EXCEL_CHART = {...}                # see CHART_TYPES and spec keys below

Workbook layout:
    「图表」   native chart that references the cells in「数据」(edit data -> chart updates)
    「数据」   the table behind the chart
    「PNG预览」 the matplotlib PNG, for side-by-side comparison
Copy the chart from Excel and paste it into PowerPoint to get an editable native PPT chart.
"""

from __future__ import annotations

import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import (AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference,
                            ScatterChart, Series)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dataio import to_number  # noqa: E402

CHART_TYPES = {
    "col": "簇状柱形图", "stacked_col": "堆积柱形图", "percent_col": "百分比堆积柱形图",
    "bar": "簇状条形图（横向）", "stacked_bar": "堆积条形图", "percent_bar": "百分比堆积条形图",
    "line": "折线图", "area": "面积图", "stacked_area": "堆积面积图",
    "pie": "饼图", "doughnut": "环形图", "scatter": "散点图",
    "heatmap": "热力表（单元格色阶，Excel 没有原生热力图）",
    "image": "仅嵌入 PNG（箱线图、小提琴图、回归线等 Excel 无原生对应的图）",
}
SPEC_KEYS = {"type", "title", "x_title", "y_title", "number_format", "data_labels", "colors",
             "secondary", "secondary_title", "secondary_number_format", "secondary_min", "secondary_max",
             "y_min", "y_max",
             "x", "y", "group", "width_cm", "height_cm"}
# matplotlib's default tab10 cycle, so the Excel chart matches the PNG
DEFAULT_COLORS = ["1F77B4", "FF7F0E", "2CA02C", "D62728", "9467BD",
                  "8C564B", "E377C2", "7F7F7F", "BCBD22", "17BECF"]
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
NOTE = "图表引用「数据」工作表中的单元格，修改数据后图表自动更新。复制图表并粘贴到 PowerPoint 即为可编辑的原生图表。"


def _labels(col: pd.Series) -> pd.Series:
    """Datetime categories -> readable labels (2024-01 for monthly data, 2024-01-05 otherwise)."""
    if pd.api.types.is_datetime64_any_dtype(col):
        fmt = "%Y-%m" if (col.dropna().dt.day == 1).all() else "%Y-%m-%d"
        return col.dt.strftime(fmt)
    if isinstance(col.dtype, pd.PeriodDtype):
        return col.astype(str)
    return col


def _cell(value):
    if value is None or (isinstance(value, float) and pd.isna(value)) or value is pd.NA or value is pd.NaT:
        return None
    if hasattr(value, "item"):  # numpy scalar
        return value.item()
    return value


def _normalize(table: pd.DataFrame, spec: dict, warnings: list[str]) -> pd.DataFrame:
    if not isinstance(table, pd.DataFrame):
        raise TypeError(f"chart_data(data) 必须返回 DataFrame，实际返回了 {type(table).__name__}。")
    table = table.reset_index() if not isinstance(table.index, pd.RangeIndex) else table.copy()
    table.columns = [str(c) for c in table.columns]
    kind = spec["type"]

    if kind == "scatter" and spec.get("group"):
        missing = [spec.get(k) for k in ("x", "y", "group") if spec.get(k) not in table.columns]
        if missing:
            raise ValueError(f"散点图分组模式需要 x/y/group 三列都在 chart_data 中，缺少 {missing}。")
        for key in ("x", "y"):
            table[spec[key]] = to_number(table[spec[key]])
        return table

    if table.shape[1] < 2:
        raise ValueError("chart_data(data) 至少需要 2 列：第 1 列是类别/x，其余列是数值系列。")
    first = table.columns[0]
    # assign by column name: iloc assignment would cast the labels back to datetime
    table[first] = to_number(table[first]) if kind == "scatter" else _labels(table[first]).astype(str)
    for col in table.columns[1:]:
        if not pd.api.types.is_numeric_dtype(table[col]):
            table[col] = to_number(table[col])
        if table[col].isna().all():
            raise ValueError(f"系列「{col}」没有任何数值，检查 chart_data 的聚合结果。")

    n = len(table)
    if kind in ("pie", "doughnut"):
        if table.shape[1] > 2:
            warnings.append(f"{CHART_TYPES[kind]}只使用第一个数值列「{table.columns[1]}」，其余列被忽略。")
        if n > 8:
            warnings.append(f"{CHART_TYPES[kind]}有 {n} 块，太多了：合并为 Top 5 +「其他」或改用条形图。")
    elif kind not in ("heatmap", "image", "scatter") and n > 60:
        warnings.append(f"Excel 图表有 {n} 个类别/时间点，会很拥挤：考虑按更粗的时间粒度汇总或保留 Top N。")
    if n > 5000:
        raise ValueError(f"chart_data 返回了 {n} 行。Excel 图表应基于汇总后的数据，而不是明细。")
    return table


def _write_table(ws, table: pd.DataFrame, number_format: str | None, start_row: int = 1) -> None:
    for j, name in enumerate(table.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=name)
        cell.font, cell.fill = Font(bold=True), HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(table.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=_cell(value))
            if number_format and j > 1 and isinstance(cell.value, (int, float)):
                cell.number_format = number_format
    for j, name in enumerate(table.columns, start=1):
        width = max([len(str(name))] + [len(str(v)) for v in table.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width * 1.6 + 2), 40)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)


def _color(spec: dict, i: int) -> str:
    colors = spec.get("colors") or DEFAULT_COLORS
    return str(colors[i % len(colors)]).lstrip("#").upper()


def _marker(color: str, size: int) -> Marker:
    marker = Marker(symbol="circle", size=size)
    marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))
    return marker


def _style_series(chart, spec: dict, offset: int = 0) -> None:
    for i, s in enumerate(chart.series):
        color = _color(spec, i + offset)
        if isinstance(chart, LineChart):
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 28575  # 2.25pt
            s.smooth = False
            s.marker = _marker(color, 5)
        elif isinstance(chart, ScatterChart):
            s.marker = _marker(color, 6)
            s.graphicalProperties.line.noFill = True
        else:
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color


def _axes(chart, spec: dict, kind: str) -> None:
    chart.x_axis.delete = False  # openpyxl >= 3.1 hides axes unless this is set
    chart.y_axis.delete = False
    chart.x_axis.title = spec.get("x_title") or None
    chart.y_axis.title = spec.get("y_title") or None
    if spec.get("number_format"):
        chart.y_axis.number_format = spec["number_format"]
        chart.y_axis.numFmt.sourceLinked = False
    if spec.get("y_min") is not None:
        chart.y_axis.scaling.min = spec["y_min"]
    if spec.get("y_max") is not None:
        chart.y_axis.scaling.max = spec["y_max"]
    chart.x_axis.tickLblPos = "low"
    if kind in ("bar", "stacked_bar", "percent_bar"):
        # keep the table order top-to-bottom and the value axis at the bottom
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"


def _build_chart(ws_data, table: pd.DataFrame, spec: dict):
    kind = spec["type"]
    n_rows, n_cols = len(table), table.shape[1]
    cats = Reference(ws_data, min_col=1, min_row=2, max_row=n_rows + 1)
    secondary = [c for c in spec.get("secondary") or [] if c in table.columns[1:]]
    primary_cols = [j for j, c in enumerate(table.columns, start=1) if j > 1 and c not in secondary]

    if kind in ("pie", "doughnut"):
        chart = PieChart() if kind == "pie" else DoughnutChart()
        chart.add_data(Reference(ws_data, min_col=2, min_row=1, max_row=n_rows + 1), titles_from_data=True)
        chart.set_categories(cats)
        for i in range(n_rows):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = _color(spec, i)
            chart.series[0].dPt.append(point)
        chart.dataLabels = DataLabelList(showPercent=True, showCatName=True, showVal=False,
                                         showSerName=False, showLeaderLines=True)
        chart.legend.position = "r"
        return chart

    if kind == "scatter":
        chart = ScatterChart()
        chart.style = 13
        if spec.get("group"):
            x_col = list(table.columns).index(spec["x"]) + 1
            y_col = list(table.columns).index(spec["y"]) + 1
            g = table[spec["group"]]
            # table was written sorted by group, so each group is a contiguous block of rows
            row = 2
            for name, size in g.value_counts(sort=False).sort_index().items():
                xs = Reference(ws_data, min_col=x_col, min_row=row, max_row=row + size - 1)
                ys = Reference(ws_data, min_col=y_col, min_row=row, max_row=row + size - 1)
                chart.series.append(Series(ys, xs, title=str(name)))
                row += size
        else:
            xs = Reference(ws_data, min_col=1, min_row=2, max_row=n_rows + 1)
            for j in range(2, n_cols + 1):
                ys = Reference(ws_data, min_col=j, min_row=1, max_row=n_rows + 1)
                chart.series.append(Series(ys, xs, title_from_data=True))
        _style_series(chart, spec)
        _axes(chart, spec, kind)
        chart.x_axis.tickLblPos = "nextTo"
        return chart

    if kind in ("line",):
        chart = LineChart()
    elif kind in ("area", "stacked_area"):
        chart = AreaChart()
        chart.grouping = "stacked" if kind == "stacked_area" else "standard"
    else:
        chart = BarChart()
        chart.type = "bar" if kind.endswith("bar") else "col"
        chart.grouping = {"stacked_col": "stacked", "stacked_bar": "stacked",
                          "percent_col": "percentStacked", "percent_bar": "percentStacked"}.get(kind, "clustered")
        if chart.grouping != "clustered":
            chart.overlap = 100
        chart.gapWidth = 60

    for j in primary_cols:
        chart.add_data(Reference(ws_data, min_col=j, min_row=1, max_row=n_rows + 1), titles_from_data=True)
    chart.set_categories(cats)
    _style_series(chart, spec)
    _axes(chart, spec, kind)

    if spec.get("data_labels"):
        chart.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                         showLegendKey=False, showPercent=False)

    if secondary:
        line = LineChart()
        for c in secondary:
            j = list(table.columns).index(c) + 1
            line.add_data(Reference(ws_data, min_col=j, min_row=1, max_row=n_rows + 1), titles_from_data=True)
        line.set_categories(cats)
        _style_series(line, spec, offset=len(primary_cols))
        line.y_axis.axId = 200
        line.y_axis.delete = False
        line.y_axis.crosses = "max"
        line.y_axis.majorGridlines = None
        line.y_axis.title = spec.get("secondary_title") or None
        if spec.get("secondary_min") is not None:
            line.y_axis.scaling.min = spec["secondary_min"]
        if spec.get("secondary_max") is not None:
            line.y_axis.scaling.max = spec["secondary_max"]
        if spec.get("secondary_number_format"):
            line.y_axis.number_format = spec["secondary_number_format"]
            line.y_axis.numFmt.sourceLinked = False
        chart += line
    return chart


def build_workbook(table: pd.DataFrame, spec: dict, out_path: str, png_path: str | None = None) -> list[str]:
    """Write the workbook; return quality warnings. Raises ValueError/TypeError on an invalid spec."""
    if not isinstance(spec, dict):
        raise TypeError("EXCEL_CHART 必须是 dict。")
    spec = dict(spec)
    kind = spec.get("type")
    if kind not in CHART_TYPES:
        raise ValueError(f"EXCEL_CHART['type'] = {kind!r} 不支持。可选: " +
                         ", ".join(f"{k}（{v}）" for k, v in CHART_TYPES.items()))
    unknown = set(spec) - SPEC_KEYS
    warnings: list[str] = [f"EXCEL_CHART 里有未识别的键 {sorted(unknown)}，已忽略。"] if unknown else []

    table = _normalize(table, spec, warnings)
    if kind == "scatter" and spec.get("group"):
        table = table.sort_values(spec["group"], kind="stable").reset_index(drop=True)

    wb = Workbook()
    ws_chart = wb.active
    ws_chart.title = "图表"
    ws_chart.page_setup.orientation = "landscape"
    ws_chart.page_setup.fitToWidth, ws_chart.page_setup.fitToHeight = 1, 0
    ws_chart.sheet_properties.pageSetUpPr.fitToPage = True
    ws_data = wb.create_sheet("数据")
    _write_table(ws_data, table, spec.get("number_format"))

    if kind == "heatmap":
        _write_table(ws_chart, table, spec.get("number_format"), start_row=3)
        ws_chart["A1"] = spec.get("title") or ""
        ws_chart["A1"].font = Font(bold=True, size=14)
        last = f"{get_column_letter(table.shape[1])}{len(table) + 3}"
        ws_chart.conditional_formatting.add(
            f"B4:{last}", ColorScaleRule(start_type="min", start_color="F7FBFF", mid_type="percentile",
                                         mid_value=50, mid_color="C6DBEF", end_type="max", end_color="4292C6"))
        ws_chart.cell(row=len(table) + 5, column=1, value="Excel 没有原生热力图，这里用单元格色阶代替；在 PPT 中可直接粘贴为表格。")
    elif kind == "image":
        if not png_path:
            raise ValueError("type='image' 需要先成功渲染 PNG。")
        warnings.append("该图表类型在 Excel 中没有原生对应，只嵌入了 PNG（不可编辑），数据仍在「数据」工作表。")
        img = XLImage(png_path)
        scale = 900 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws_chart.add_image(img, "A1")
    else:
        chart = _build_chart(ws_data, table, spec)
        chart.title = spec.get("title") or None
        chart.width = spec.get("width_cm", 25.4)   # 16:9, fits a PPT slide body
        chart.height = spec.get("height_cm", 14.3)
        if table.shape[1] <= 2 and kind not in ("pie", "doughnut"):
            chart.legend = None
        elif chart.legend is not None and kind not in ("pie", "doughnut"):
            chart.legend.position = "b"
        ws_chart.add_chart(chart, "A1")
        ws_chart.cell(row=31, column=1, value=NOTE).font = Font(color="808080")

    if png_path and os.path.exists(png_path) and kind != "image":
        ws_png = wb.create_sheet("PNG预览")
        img = XLImage(png_path)
        scale = 900 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws_png.add_image(img, "A1")

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return warnings
